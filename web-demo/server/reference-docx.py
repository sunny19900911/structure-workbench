"""Populate the retained Word patterns without reconstructing tables from HTML."""
import json,sys,base64,io,re
from copy import deepcopy
from pathlib import Path
from lxml import etree as E
from docx import Document
from docx.shared import Pt
ROOT=Path(__file__).resolve().parents[1]
PAT=json.loads((ROOT/'design/reference-expansion/native-tables.json').read_text(encoding='utf-8'))
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';N={'w':W};q=lambda s:'{'+W+'}'+s
def text_into(p,text):
 rp=p.find('w:r/w:rPr',N)
 for c in list(p):
  if c.tag!=q('pPr'):p.remove(c)
 r=E.SubElement(p,q('r'))
 if rp is not None:r.append(deepcopy(rp))
 for i,line in enumerate(str(text or '').split('\n')):
  if i:E.SubElement(r,q('br'))
  t=E.SubElement(r,q('t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=line
def build(spec,dest):
 doc=Document(ROOT/'design/reference-expansion/base.docx');body=doc._element.body
 def append(n):body.insert(len(body)-1,n)
 for block in spec['blocks']:
  kind=block.get('kind')
  if kind=='table':
   key=str(block['template']);original=E.fromstring(PAT['tables'][key]);table=deepcopy(original)
   original_rows=original.findall('w:tr',N)
   for r in table.findall('w:tr',N):table.remove(r)
   for item in block['rows']:
    r=deepcopy(original_rows[item['source']]);cells=r.findall('w:tc',N)
    if len(item['values'])!=len(cells):raise ValueError('Table cell count mismatch')
    for c,value in zip(cells,item['values']):
     ps=c.findall('w:p',N);p=ps[0] if ps else E.SubElement(c,q('p'))
     for n in list(c):
      if n not in [c.find('w:tcPr',N),p]:c.remove(n)
     text_into(p,value)
    table.append(r)
   append(table)
  elif kind=='image':
   raw=block.get('data','')
   if not re.fullmatch(r'data:image/(?:png|jpeg);base64,[A-Za-z0-9+/=]+',raw):raise ValueError('Invalid image')
   data=base64.b64decode(raw.split(',',1)[1],validate=True)
   if len(data)>5*1024*1024:raise ValueError('Image too large')
   p=doc.add_paragraph();p.alignment=1
   shape=p.add_run().add_picture(io.BytesIO(data),width=Pt(450))
   if shape.height>Pt(340):shape.width=int(shape.width*Pt(340)/shape.height);shape.height=Pt(340)
   p.paragraph_format.keep_with_next=True
  elif kind in PAT['paragraphs']:
   p=E.fromstring(PAT['paragraphs'][kind]);text_into(p,block.get('text',''));append(p)
  else:raise ValueError('Unknown document block')
 # Native PAGE field, matching the reference's small right-aligned page number.
 footer=doc.sections[0].footer.paragraphs[0];footer.alignment=2
 run=footer.add_run();run.font.size=Pt(8)
 field=E.SubElement(run._r,q('fldChar'));field.set(q('fldCharType'),'begin')
 instruction=E.SubElement(run._r,q('instrText'));instruction.text=' PAGE '
 end=E.SubElement(run._r,q('fldChar'));end.set(q('fldCharType'),'end')
 doc.save(dest)
if __name__=='__main__':build(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8')),sys.argv[2])
