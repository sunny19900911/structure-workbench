from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import openpyxl

root = Path(__file__).resolve().parents[1]
old = openpyxl.load_workbook(root / 'outputs/主楼钢筋用量_按构件汇总.xlsx', data_only=True)
out = root / 'outputs/20260927-chart-refresh/主楼钢筋用量_图表排版优化.xlsx'
new = openpyxl.load_workbook(out, data_only=True)
formulas = openpyxl.load_workbook(out, data_only=False)
for row in range(6, 20):
    area=old['按楼层汇总'].cell(row,2).value
    assert new['按楼层汇总'].cell(row,2).value==area
    for col in range(3,7):
        original_mass=old['按楼层汇总'].cell(row,col).value*1000
        expected=original_mass/area
        actual=new['按楼层汇总'].cell(row,col).value
        assert abs(expected-actual)<1e-8,(row,col,expected,actual)
        if row<19:
            assert abs(new['工程量来源'].cell(row,col).value-original_mass)<1e-8
    assert abs(new['按楼层汇总'].cell(row,7).value-old['按楼层汇总'].cell(row,8).value)<1e-8
    assert abs(new['按楼层汇总'].cell(row,8).value-old['按楼层汇总'].cell(row,7).value*1000)<1e-6
assert formulas['按楼层汇总']['G19'].value == '=IF(B19>0,H19/B19,"待补面积")'
assert formulas['按楼层汇总'].freeze_panes == 'A6'
assert abs(new['用量总览']['B9'].value-new['按楼层汇总']['G19'].value)<1e-8
assert abs(new['用量总览']['B11'].value-new['按楼层汇总']['G19'].value)<1e-10
for row,col in [(5,3),(6,4),(7,5),(8,6)]:
    assert abs(new['用量总览'].cell(row,2).value-new['按楼层汇总'].cell(19,col).value)<1e-10
with ZipFile(out) as archive:
    assert archive.testzip() is None
    charts=[name for name in archive.namelist() if '/charts/chart' in name and name.endswith('.xml')]
    assert len(charts)==2
    ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    for chart in charts:
        node=ET.fromstring(archive.read(chart))
        refs=[el.text for el in node.findall('.//c:f',ns)]
        assert refs, chart
        print(chart, refs)
print('Verified: mass preserved; component kg/m2 per floor and weighted totals, summary links, frozen headings, ZIP and 2 native charts.')
